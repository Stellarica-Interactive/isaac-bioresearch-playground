"""Importer for Cook et al. 2019 (Nature 571:63-71), hermaphrodite.

The whole-animal *C. elegans* connectome: 302 neurons, 135 muscles and 36 other
cells, reconstructed from electron micrographs of an adult hermaphrodite.

Why this dataset matters here
-----------------------------

Unlike Witvliet et al. 2021, this reconstruction covers the **entire animal**,
including the ventral nerve cord motor neurons (DA, DB, VA, VB, VC, DD, VD, AS)
and all 95 body wall muscles. The travelling wave that makes the worm crawl is
generated in the ventral cord, so this is the first dataset in the repository
that could support a locomotion model at all.

Format: adjacency matrices, not an edge list
--------------------------------------------

A genuinely different shape from Witvliet, hence a separate importer. Each sheet
is a matrix with presynaptic cells down column C and postsynaptic cells across
row 3, data starting at row 4 / column D. Empty cells mean "no connection", not
"zero connections", and are skipped.

The gap junction sheets, and why the published totals are odd
------------------------------------------------------------

There are two, and the naming is misleading:

``hermaphrodite gap jn symmetric``
    The full matrix, mirrored: each junction appears in both directions. 2883
    non-zero entries summing to 23313. These are the numbers quoted in published
    summaries.

``hermaphrodite gap jn asymmetric``
    **Not** asymmetric conductances — this is the *upper-triangle* form, with each
    junction written exactly once. 1450 entries summing to 11680.

The second form is precisely our storage convention, so that is the sheet we read.
It removes any chance of double counting at import time. The published symmetric
figures are then reproduced by :meth:`Connectome.electrical_directed_view`, and the
arithmetic explains why they are odd numbers::

    entries: 2 * (1450 - 17) + 17 = 2883
    weight:  2 *  11680      - 47 = 23313

The 17 diagonal entries are gap junctions from a cell to itself, carrying a total
weight of 47. A self-junction lies on the matrix diagonal and is written once, not
twice, which is what makes both totals odd. The same phenomenon appears in every
Witvliet dataset; see ``docs/model_assumptions.md``.

What this data still does not contain
-------------------------------------

Synaptic sign, synaptic strength in physiological units, neurotransmitter
identity, functional role, or any neural activity. As with every connectome.
"""

from __future__ import annotations

from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from pathlib import Path

from common.data.registry import ConnectomeImporter, SourceFile, register_importer
from common.data.schemas import (
    Cell,
    CellCategory,
    Connection,
    Connectome,
    Provenance,
    Scope,
    SynapseType,
)
from worm.importers.naming import canonical_cell_id, category_of
from worm.importers.sources import source_file

FILE_KEY = "connectome"
DATASET_ID = "cook_2019_herm"

#: Cell category is not measured by the connectome; it comes from the cell registry.
REGISTRY_SOURCE_ID = "wormatlas_cook_2019_via_cect"

CHEMICAL_SHEET = "hermaphrodite chemical"
#: The upper-triangle sheet: each gap junction exactly once. See the module docstring.
ELECTRICAL_SHEET = "hermaphrodite gap jn asymmetric"
#: Read only to cross-check the above; never stored.
ELECTRICAL_MIRROR_SHEET = "hermaphrodite gap jn symmetric"

#: Zero-based positions in every matrix sheet.
HEADER_ROW = 2
NAME_COL = 2
FIRST_DATA_ROW = 3
FIRST_DATA_COL = 3


@dataclass(frozen=True, slots=True)
class MatrixEntry:
    pre: str
    post: str
    weight: int


def _read_matrix(rows: Sequence[Sequence[object]], sheet: str) -> list[MatrixEntry]:
    """Parse one adjacency-matrix sheet into raw (pre, post, weight) triples.

    Cell labels are returned exactly as written; canonicalization happens in the
    caller so that both axes and both sheets go through the same single path.
    """
    if len(rows) <= FIRST_DATA_ROW:
        raise ValueError(f"{sheet}: sheet has no data rows")

    header = rows[HEADER_ROW]
    post_cols = [
        (j, str(v).strip())
        for j, v in enumerate(header)
        if j >= FIRST_DATA_COL and v is not None and str(v).strip()
    ]
    pre_rows = [
        (i, str(r[NAME_COL]).strip())
        for i, r in enumerate(rows)
        if i >= FIRST_DATA_ROW
        and len(r) > NAME_COL
        and r[NAME_COL] is not None
        and str(r[NAME_COL]).strip()
    ]
    if not post_cols or not pre_rows:
        raise ValueError(
            f"{sheet}: found {len(pre_rows)} presynaptic rows and {len(post_cols)} "
            "postsynaptic columns. The upstream layout changed; do not adjust the "
            "offsets without re-verifying against worm/data/reference_totals.toml."
        )

    out: list[MatrixEntry] = []
    for i, pre in pre_rows:
        row = rows[i]
        for j, post in post_cols:
            if j >= len(row):
                continue
            value = row[j]
            if value is None or value == "":
                continue  # absent, not zero
            if isinstance(value, str) and not value.strip():
                continue
            try:
                weight = int(float(value))  # type: ignore[arg-type]
            except (TypeError, ValueError):
                raise ValueError(
                    f"{sheet}: non-numeric cell {value!r} at {pre}->{post}"
                ) from None
            if weight == 0:
                continue
            if weight < 0:
                raise ValueError(f"{sheet}: negative weight {weight} at {pre}->{post}")
            out.append(MatrixEntry(pre, post, weight))
    return out


@dataclass(frozen=True)
class Cook2019HermImporter:
    """Parses the hermaphrodite sheets of Cook et al. 2019 Supplementary Information 5."""

    @property
    def dataset_id(self) -> str:
        return DATASET_ID

    def source_files(self) -> list[SourceFile]:
        return [source_file(DATASET_ID, key=FILE_KEY)]

    def parse(self, files: Mapping[str, Path]) -> Connectome:
        from openpyxl import load_workbook

        path = files[FILE_KEY]
        spec = self.source_files()[0]

        workbook = load_workbook(path, read_only=True, data_only=True)
        missing = [
            s
            for s in (CHEMICAL_SHEET, ELECTRICAL_SHEET, ELECTRICAL_MIRROR_SHEET)
            if s not in workbook.sheetnames
        ]
        if missing:
            raise ValueError(f"{path.name}: missing expected sheet(s) {missing}")

        chemical_raw = _read_matrix(
            list(workbook[CHEMICAL_SHEET].iter_rows(values_only=True)), CHEMICAL_SHEET
        )
        electrical_raw = _read_matrix(
            list(workbook[ELECTRICAL_SHEET].iter_rows(values_only=True)), ELECTRICAL_SHEET
        )
        mirror_raw = _read_matrix(
            list(workbook[ELECTRICAL_MIRROR_SHEET].iter_rows(values_only=True)),
            ELECTRICAL_MIRROR_SHEET,
        )
        workbook.close()

        _check_mirror_consistency(electrical_raw, mirror_raw, path)

        raw_labels: dict[str, set[str]] = {}

        def canonical(label: str) -> str:
            cell_id = canonical_cell_id(label)
            raw_labels.setdefault(cell_id, set()).add(label)
            return cell_id

        connections: list[Connection] = []
        seen: set[tuple[str, str, SynapseType]] = set()

        for entry in chemical_raw:
            pre, post = canonical(entry.pre), canonical(entry.post)
            _add(connections, seen, pre, post, SynapseType.CHEMICAL, entry.weight, path)

        for entry in electrical_raw:
            pre, post = canonical(entry.pre), canonical(entry.post)
            if pre > post:
                # Gap junctions are undirected; store the canonical orientation.
                pre, post = post, pre
            _add(connections, seen, pre, post, SynapseType.ELECTRICAL, entry.weight, path)

        provenance = Provenance(
            source_id=DATASET_ID,
            kind="connectome",
            citation=spec.citation,
            url=spec.url,
            license=spec.license,
            doi=spec.doi,
            file_name=spec.filename,
            sha256=spec.sha256,
            notes=(
                "Whole-animal reconstruction: includes the ventral nerve cord motor "
                "neurons and all 95 body wall muscles. Gap junctions read from the "
                "upper-triangle 'asymmetric' sheet, which stores each junction once."
            ),
        )
        registry_spec = source_file(REGISTRY_SOURCE_ID)
        registry_provenance = Provenance(
            source_id=REGISTRY_SOURCE_ID,
            kind="annotation",
            citation=registry_spec.citation,
            url=registry_spec.url,
            license=registry_spec.license,
            doi=registry_spec.doi,
            notes="Cell category (neuron / muscle / glia / other) via worm/data/cells.csv.",
        )

        cells = tuple(
            Cell(
                id=cell_id,
                category=_category(cell_id, path),
                source_labels=tuple(sorted(labels)),
                field_sources={"category": REGISTRY_SOURCE_ID},
            )
            for cell_id, labels in sorted(raw_labels.items())
        )

        return Connectome(
            id=DATASET_ID,
            organism="Caenorhabditis elegans",
            sex="hermaphrodite",
            stage="adult",
            scope=Scope.WHOLE_ANIMAL,
            cells=cells,
            connections=tuple(connections),
            provenance=provenance,
            sources={REGISTRY_SOURCE_ID: registry_provenance},
        )


def _add(
    connections: list[Connection],
    seen: set[tuple[str, str, SynapseType]],
    pre: str,
    post: str,
    synapse_type: SynapseType,
    weight: int,
    path: Path,
) -> None:
    key = (pre, post, synapse_type)
    if key in seen:
        raise ValueError(
            f"{path.name}: duplicate connection {pre}->{post} ({synapse_type}) after "
            "canonicalization. Two source labels collapsed onto one cell; merging them "
            "would silently sum their synapse counts, so this is a hard error."
        )
    seen.add(key)
    connections.append(
        Connection(pre=pre, post=post, synapse_type=synapse_type, weight=weight)
    )


def _check_mirror_consistency(
    upper: Sequence[MatrixEntry], mirrored: Sequence[MatrixEntry], path: Path
) -> None:
    """Verify the two gap-junction sheets describe the same network.

    The workbook ships both a once-per-pair sheet and a fully mirrored one. They are
    redundant, which makes them a free correctness check on our reading of the layout:
    if our offsets or our orientation were wrong, these would disagree.
    """
    mirror = {(e.pre, e.post): e.weight for e in mirrored}
    problems: list[str] = []
    for e in upper:
        for key in ((e.pre, e.post), (e.post, e.pre)):
            if mirror.get(key) != e.weight:
                problems.append(f"{e.pre}<->{e.post}: {e.weight} vs mirror {mirror.get(key)}")
                break
        if len(problems) >= 5:
            break

    expected_entries = 2 * sum(1 for e in upper if e.pre != e.post) + sum(
        1 for e in upper if e.pre == e.post
    )
    if not problems and len(mirrored) != expected_entries:
        problems.append(
            f"mirrored sheet has {len(mirrored)} entries, expected {expected_entries} "
            f"from {len(upper)} stored junctions"
        )

    if problems:
        raise ValueError(
            f"{path.name}: the '{ELECTRICAL_SHEET}' and '{ELECTRICAL_MIRROR_SHEET}' "
            "sheets disagree, so our reading of one of them is wrong:\n  "
            + "\n  ".join(problems)
        )


def _category(cell_id: str, path: Path) -> CellCategory:
    try:
        return category_of(cell_id)
    except KeyError as exc:
        raise ValueError(f"{path.name}: {exc}") from None


IMPORTER = register_importer(Cook2019HermImporter())

# Static type check: the dataclass really does satisfy the protocol.
_: ConnectomeImporter = IMPORTER
