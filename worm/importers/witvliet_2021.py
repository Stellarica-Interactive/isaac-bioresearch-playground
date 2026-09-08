"""Importer for Witvliet et al. 2021 (Nature 596:257-261).

Eight individual *C. elegans* brains reconstructed by serial-section electron
microscopy, spanning birth (L1) to adulthood. Datasets 7 and 8 are adults.

What this data is
-----------------

Each row of the source spreadsheet is one *connection* between two cells, with a
weight that is the number of presynaptic active zones agreed upon by at least two
of three independent human annotators. So the weight is a count of physical
structures visible in the micrographs — not a conductance, not a strength, and
not a rate.

What this data is **not**
-------------------------

**It is not a whole-animal connectome.** Witvliet et al. reconstructed the brain:
the nerve ring and the ganglia immediately around it. There are no ventral-cord
motor neurons (DA, DB, VA, VB, VC, VD, DD, AS) and only the anterior eight body
wall muscle segments of each quadrant. The circuitry that generates the worm's
undulating crawling gait lives in the ventral cord, so **a locomotion model
cannot be built from these datasets**; use a whole-animal reconstruction for
that. :attr:`~common.data.schemas.Scope.HEAD` records this in the data itself so
downstream code can refuse rather than silently produce a worm that cannot move.

It also carries no synaptic sign, no neurotransmitter identity, no functional
role, and no information about the animal's neural activity while it was alive.
Those are attached separately by :mod:`common.data.overlay`, or are simply absent.

Format
------

One sheet, header row 1, four columns: ``pre``, ``post``, ``type``
(``"chemical"`` or ``"electrical"``), ``synapses`` (integer count).

Gap junctions appear **once** per pair, in an arbitrary direction, which suits
our storage convention directly. They include self-junctions — a gap junction
between two processes of the same cell — in every one of the eight datasets.
Those are preserved as stored self-loops; see ``docs/model_assumptions.md``.
"""

from __future__ import annotations

from collections.abc import Mapping
from dataclasses import dataclass
from pathlib import Path

from common.data.registry import ConnectomeImporter, SourceFile, register_importer
from common.data.schemas import (
    Cell,
    CellCategory,
    Connection,
    Connectome,
    Provenance,
    Scope,
    SynapseType,
)
from worm.importers.naming import canonical_cell_id, category_of
from worm.importers.sources import source_file

FILE_KEY = "connectome"

#: Cell category is not measured by the connectome; it comes from the cell registry.
REGISTRY_SOURCE_ID = "wormatlas_cook_2019_via_cect"

_STAGES = {1: "L1", 2: "L1", 3: "L1", 4: "L1", 5: "L2", 6: "L3", 7: "adult", 8: "adult"}


@dataclass(frozen=True)
class WitvlietImporter:
    """Parses one of the eight Witvliet datasets."""

    number: int

    @property
    def dataset_id(self) -> str:
        return f"witvliet_2021_{self.number}"

    @property
    def stage(self) -> str:
        return _STAGES[self.number]

    def source_files(self) -> list[SourceFile]:
        return [source_file(self.dataset_id, key=FILE_KEY)]

    def parse(self, files: Mapping[str, Path]) -> Connectome:
        from openpyxl import load_workbook

        path = files[FILE_KEY]
        spec = self.source_files()[0]

        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.worksheets[0]
        rows = sheet.iter_rows(values_only=True)

        header = next(rows, None)
        expected = ("pre", "post", "type", "synapses")
        if tuple(str(h).strip() for h in (header or ())) != expected:
            raise ValueError(
                f"{path.name}: expected header {expected}, found {header}. "
                "The upstream file layout changed; do not adjust column indices "
                "without re-verifying against worm/data/reference_totals.toml."
            )

        raw_labels: dict[str, set[str]] = {}
        connections: list[Connection] = []
        seen: set[tuple[str, str, SynapseType]] = set()

        for line_no, row in enumerate(rows, start=2):
            if row is None or row[0] is None:
                continue
            pre_raw, post_raw, type_raw, weight_raw = row[0], row[1], row[2], row[3]

            pre = canonical_cell_id(str(pre_raw))
            post = canonical_cell_id(str(post_raw))
            raw_labels.setdefault(pre, set()).add(str(pre_raw))
            raw_labels.setdefault(post, set()).add(str(post_raw))

            synapse_type = _parse_synapse_type(str(type_raw), path, line_no)
            weight = int(weight_raw)
            if weight < 1:
                raise ValueError(
                    f"{path.name}:{line_no}: connection {pre}->{post} has weight {weight}. "
                    "A connection that was not observed should be absent from the file."
                )

            if synapse_type is SynapseType.ELECTRICAL and pre > post:
                # Gap junctions are undirected; store the canonical orientation.
                pre, post = post, pre

            key = (pre, post, synapse_type)
            if key in seen:
                raise ValueError(
                    f"{path.name}:{line_no}: duplicate connection {pre}->{post} "
                    f"({synapse_type}). Merging duplicates would silently change weights, "
                    "so this is a hard error."
                )
            seen.add(key)
            connections.append(
                Connection(pre=pre, post=post, synapse_type=synapse_type, weight=weight)
            )

        workbook.close()

        provenance = Provenance(
            source_id=self.dataset_id,
            kind="connectome",
            citation=spec.citation,
            url=spec.url,
            license=spec.license,
            doi=spec.doi,
            file_name=spec.filename,
            sha256=spec.sha256,
            notes=(
                "Brain-only reconstruction: no ventral-cord motor neurons, "
                "body wall muscles limited to the anterior eight segments per quadrant."
            ),
        )

        # The connectome tells us that a cell exists and what it connects to. It does
        # not tell us whether that cell is a neuron, a muscle or a glial cell — that
        # classification comes from the cell registry, so it is cited accordingly.
        registry_spec = source_file(REGISTRY_SOURCE_ID)
        registry_provenance = Provenance(
            source_id=REGISTRY_SOURCE_ID,
            kind="annotation",
            citation=registry_spec.citation,
            url=registry_spec.url,
            license=registry_spec.license,
            doi=registry_spec.doi,
            notes="Cell category (neuron / muscle / glia / other) via worm/data/cells.csv.",
        )

        cells = tuple(
            Cell(
                id=cell_id,
                category=_category(cell_id, path),
                source_labels=tuple(sorted(labels)),
                field_sources={"category": REGISTRY_SOURCE_ID},
            )
            for cell_id, labels in sorted(raw_labels.items())
        )

        return Connectome(
            id=self.dataset_id,
            organism="Caenorhabditis elegans",
            sex="hermaphrodite",
            stage=self.stage,
            scope=Scope.HEAD,
            cells=cells,
            connections=tuple(connections),
            provenance=provenance,
            sources={REGISTRY_SOURCE_ID: registry_provenance},
        )


def _parse_synapse_type(raw: str, path: Path, line_no: int) -> SynapseType:
    value = raw.strip().lower()
    if value == "chemical":
        return SynapseType.CHEMICAL
    if value == "electrical":
        return SynapseType.ELECTRICAL
    raise ValueError(
        f"{path.name}:{line_no}: unknown synapse type {raw!r}; expected "
        "'chemical' or 'electrical'"
    )


def _category(cell_id: str, path: Path) -> CellCategory:
    try:
        return category_of(cell_id)
    except KeyError as exc:
        raise ValueError(f"{path.name}: {exc}") from None


IMPORTERS = [register_importer(WitvlietImporter(n)) for n in range(1, 9)]

# Static type check: the dataclass really does satisfy the protocol.
_: ConnectomeImporter = IMPORTERS[0]
